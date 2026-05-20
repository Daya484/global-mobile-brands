"""
Cloud Run Job: Generate Distributor Excel files and upload to GCS
===============================================================

Goal:
- Generate Excel files per country and per distributor.
- Upload files to GCS in the exact structure:
    gs://<bucket_name>/<landing_folder>/<COUNTRY_FOLDER>/<DIST_CODE>_<YYYYMMDD>.xlsx

Example:
    gs://mobile_brands/landing/AUSTRALIA/AUSDST01_20260509.xlsx

Airflow integration (optional):
- Pass RUN_DATE={{ ds }} so date changes daily.
- If RUN_DATE is not provided, job uses today's date.
"""

import os
import sys
import json
import random
import logging
from datetime import date

import openpyxl
from openpyxl.styles import Font
from google.cloud import storage


# -----------------------------------------------------------------------------
# 1) CONFIG LOADING
# -----------------------------------------------------------------------------

def load_config() -> dict:
    """
    Loads config from config/config.json.
    Keeping parameters in JSON helps new members adjust behavior without code changes.
    """
    with open("config/config.json", "r") as f:
        return json.load(f)


def setup_logging(level: str) -> logging.Logger:
    """
    Configure logging to stdout so Cloud Run / Cloud Logging can capture it.
    """
    logging.basicConfig(
        level=getattr(logging, level.upper(), logging.INFO),
        format="%(asctime)s [%(levelname)s] %(message)s",
        stream=sys.stdout,
    )
    return logging.getLogger("excel_generator")


# -----------------------------------------------------------------------------
# 2) DATA CATALOG (Base catalogue in INR)
#    This is the product master used for all countries (price converted using fx + factor).
# -----------------------------------------------------------------------------

CATALOGUE_INR = {
    "Samsung": [
        ("Galaxy A56 5G", 38990, "1029384756102938"),
        ("Galaxy M17 5G", 13965, "1928374650192837"),
        ("Galaxy S26 Ultra", 121870, "1938475620193847"),
        ("Galaxy F07", 9499, "2039485710293847"),
        ("Galaxy S26", 76500, "2948571029384756"),
        ("Galaxy F34 5G", 21999, "3847562910384756"),
        ("Galaxy S24 Ultra 5G", 134999, "4729104857291038"),
        ("Galaxy A36 5G", 35208, "4857291038475629"),
        ("Galaxy M35 5G", 17499, "5729103847562910"),
        ("Galaxy S22 5G", 85999, "5829103847562910"),
        ("Galaxy M56 5G", 23499, "7584930218475629"),
        ("Galaxy Z Fold 5", 159999, "8392048175930284"),
        ("Galaxy A07", 10300, "8475629103847562"),
        ("Galaxy S24 FE", 39999, "9283746501928374"),
        ("Galaxy A06", 7999, "9384756201938475"),
    ],
    "Vivo": [
        ("X Fold3 Pro", 159999, "3829104857291030"),
        ("X100 Pro", 89999, "5729104857291031"),
        ("X80 Pro", 79999, "1938475620193841"),
        ("X90 Pro", 74999, "5829103847562911"),
        ("X100", 63999, "2948571029384751"),
        ("V30 Pro", 41999, "9283746501928371"),
        ("V30", 33999, "1029384756102931"),
        ("V29 5G", 32999, "4857291038475621"),
        ("T2 Pro 5G", 23999, "7584930218475621"),
        ("Y200 5G", 21999, "3847562910384751"),
        ("Y56 5G", 16999, "5729103847562911"),
        ("Y28 5G", 13999, "1928374650192831"),
        ("T2x 5G", 12999, "8475629103847561"),
        ("Y16", 10499, "2039485710293841"),
        ("Y02t", 8999, "9384756201938471"),
    ],
    "Oppo": [
        ("Find N3 Fold", 149999, "8392048175930281"),
        ("Find X7 Ultra", 99999, "4729104857291031"),
        ("Find N2 Flip", 89999, "1938475620193841"),
        ("Find X6 Pro", 84999, "5829103847562911"),
        ("Reno 10 Pro+ 5G", 54999, "2948571029384751"),
        ("Reno 11 Pro 5G", 39999, "9283746501928371"),
        ("Reno 10 5G", 32999, "1029384756102931"),
        ("Reno 11 5G", 29999, "4857291038475621"),
        ("F23 5G", 24999, "7584930218475621"),
        ("F25 Pro 5G", 23999, "3847562910384751"),
        ("A79 5G", 19999, "5729103847562911"),
        ("A59 5G", 14999, "1928374650192831"),
        ("A58", 13999, "8475629103847561"),
        ("A38", 12999, "2039485710293841"),
        ("A18", 9999, "9384756201938471"),
    ],
    "OnePlus": [
        ("OnePlus Open", 139999, "8392048175930280"),
        ("OnePlus 10 Pro", 66999, "4729104857291030"),
        ("OnePlus 12", 64999, "1938475620193840"),
        ("OnePlus 9 Pro", 64999, "5829103847562910"),
        ("OnePlus 11 5G", 56999, "2948571029384750"),
        ("OnePlus 12R", 39999, "9283746501928370"),
        ("OnePlus 11R", 35999, "1029384756102930"),
        ("OnePlus Nord 3", 33999, "4857291038475620"),
        ("OnePlus Nord 4", 29999, "7584930218475620"),
        ("OnePlus Nord CE 4", 24999, "3847562910384750"),
    ],
    "Apple": [
        ("iPhone 16 Pro Max", 159900, "8392048175930283"),
        ("iPhone 16 Pro", 134900, "4729104857291033"),
        ("iPhone 16", 90000, "1938475620193843"),
        ("iPhone 15 Pro Max", 139900, "5829103847562913"),
        ("iPhone 15 Pro", 114900, "2948571029384753"),
        ("iPhone 15 Plus", 89900, "3847562910384753"),
        ("iPhone 15", 79900, "9283746501928373"),
        ("iPhone 14 Pro Max", 129900, "1029384756102933"),
        ("iPhone 14 Pro", 99000, "4857291038475623"),
        ("iPhone 14", 60000, "7584930218475623"),
    ],
}


# -----------------------------------------------------------------------------
# 3) EXCEL OUTPUT SCHEMA (Columns)
# -----------------------------------------------------------------------------

HEADERS = [
    "Brand", "Model", "Price", "Distributor_code",
    "Retailer_code", "Store_code", "EAN_code",
    "Currency", "Stock_units", "Sale_units"
]


# -----------------------------------------------------------------------------
# 4) HELPERS (Price conversion + random stock generation)
# -----------------------------------------------------------------------------

def local_price(price_inr: int, fx: float, factor: float) -> int:
    """Convert INR price to local currency using fx and additional factor."""
    return max(1, int(round(price_inr * fx * factor)))


def rand_stock(price: int, scale: float) -> tuple[int, int]:
    """
    Generate realistic stock/sales values based on price bucket.
    Higher-priced phones typically have lower stock.
    """
    if price >= 100000:
        s_min, s_max = 3, 20
    elif price >= 50000:
        s_min, s_max = 5, 35
    elif price >= 25000:
        s_min, s_max = 10, 60
    elif price >= 15000:
        s_min, s_max = 20, 90
    else:
        s_min, s_max = 40, 180

    stock = random.randint(int(s_min * scale), int(s_max * scale))
    sales = random.randint(1, stock)
    return stock, sales


# -----------------------------------------------------------------------------
# 5) MAIN GENERATION + UPLOAD
# -----------------------------------------------------------------------------

def main():
    cfg = load_config()
    log = setup_logging(cfg.get("log_level", "INFO"))

    # Allow Airflow to control "today" using RUN_DATE={{ ds }}
    # If RUN_DATE is not passed, default to system date
    run_date = os.getenv("RUN_DATE")
    today = run_date.replace("-", "") if run_date else date.today().strftime("%Y%m%d")

    # Basic config
    env = os.getenv("ENV", "dv")

    # ✅ Your project naming standard
    project_name = f"{env}-env"

    # ✅ Bucket name from config
    bucket_name = os.getenv("BUCKET_NAME",f"{cfg['bucket_base']}_{env}")
    landing_folder = os.getenv("LANDING_FOLDER", cfg["landing_folder"])
    tmp_dir = cfg.get("tmp_dir", "/tmp/distributor_files")
    os.makedirs(tmp_dir, exist_ok=True)

    num_distributors = int(os.getenv("NUM_DISTRIBUTORS", cfg.get("num_distributors", 15)))
    num_retailers = int(os.getenv("NUM_RETAILERS", cfg.get("num_retailers", 10)))
    brand_order = cfg.get("brand_order", list(CATALOGUE_INR.keys()))
    countries = cfg["countries"]

    # Optional deterministic randomness for repeatable output
    seed = cfg.get("random_seed")
    if seed is not None:
        random.seed(seed)

    log.info("Start Excel generation | bucket=%s | landing=%s | date=%s | distributors=%s | retailers=%s",
             bucket_name, landing_folder, today, num_distributors, num_retailers)

    storage_client = storage.Client()
    bucket = storage_client.bucket(bucket_name)

    for country_name, c in countries.items():
        log.info("Country=%s (folder=%s, prefix=%s)", country_name, c["gcs_folder"], c["prefix"])

        for d in range(1, num_distributors + 1):
            dist_code = f"{c['prefix']}DST{d:02d}"
            scale = 0.85 + (d - 1) * 0.03

            # Build a retailer->stores hierarchy
            hierarchy = {}
            for r in range(1, num_retailers + 1):
                retailer_code = f"{dist_code}RET{r:02d}"
                hierarchy[retailer_code] = [
                    f"{retailer_code}ST{s:02d}"
                    for s in range(1, random.randint(5, 10) + 1)
                ]

            # Create workbook (one file per distributor)
            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            # Create brand sheets in a fixed order so files are consistent
            for brand in brand_order:
                ws = wb.create_sheet(brand)

                # Write header row with bold formatting
                for col, header in enumerate(HEADERS, 1):
                    ws.cell(row=1, column=col, value=header).font = Font(bold=True)

                row = 2
                for retailer_code, stores in hierarchy.items():
                    for store_code in stores:
                        for model, price_inr, ean in CATALOGUE_INR[brand]:
                            price = local_price(price_inr, c["fx"], c["factor"])
                            stock_units, sale_units = rand_stock(price, scale)

                            values = [
                                brand, model, price, dist_code,
                                retailer_code, store_code, ean,
                                c["currency"], stock_units, sale_units
                            ]

                            for col, v in enumerate(values, 1):
                                ws.cell(row=row, column=col, value=v)
                            row += 1

            # File naming rule: <DIST_CODE>_<YYYYMMDD>.xlsx
            file_name = f"{dist_code}_{today}.xlsx"
            local_path = os.path.join(tmp_dir, file_name)

            # Save Excel file to local temp directory (Cloud Run allows /tmp)
            wb.save(local_path)

            # Upload to GCS path: landing/<COUNTRY_FOLDER>/<filename>
            gcs_path = f"{landing_folder}/{c['gcs_folder']}/{file_name}"

            blob = bucket.blob(gcs_path)
            blob.metadata = {"run_date": today, "country": c["gcs_folder"], "distributor": dist_code}
            blob.upload_from_filename(local_path)

            # Clean up local file to avoid /tmp filling up
            try:
                os.remove(local_path)
            except Exception:
                pass

            log.info("Uploaded → gs://%s/%s", bucket_name, gcs_path)

    log.info("✅ All files generated successfully.")


if __name__ == "__main__":
    main()